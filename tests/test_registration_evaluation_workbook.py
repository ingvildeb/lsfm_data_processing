from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from lsfm_data_processing.registration_and_transforms.project_workflow.models import (
    RescueRequest,
    RescueRequestKey,
    RescueRequestRecord,
    RescueRequestStatus,
    RescueStrategy,
    RunEvaluation,
)
from lsfm_data_processing.registration_and_transforms.project_workflow.schema import (
    RESCUE_REQUESTS_SHEET,
    RESCUE_REQUEST_COLUMNS,
    RUN_EVALUATION_COLUMNS,
    RUN_EVALUATIONS_SHEET,
    WORKFLOW_LISTS_SHEET,
)
from lsfm_data_processing.registration_and_transforms.project_workflow.rescue_catalog import (
    canonical_rescue_strategies,
    expected_rescue_run_paths,
)
from lsfm_data_processing.registration_and_transforms.project_workflow.workbook import (
    apply_evaluation_workbook_plan,
    build_evaluation_workbook_plan,
    read_evaluation_workbook,
)


def _write_legacy_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = RUN_EVALUATIONS_SHEET
    sheet.append(
        (
            "Subject ID",
            "Run path",
            "Preset name",
            "Manifest status",
            "Result",
            "Comments",
            "Selected",
            "Evaluator",
            "Evaluation date",
            "Proposed rescue",
        )
    )
    sheet.append(
        (
            "A",
            r"registration_runs\baseline",
            "baseline",
            "planned",
            "minor_localized",
            "Keep this note",
            "yes",
            "Evaluator",
            date(2026, 9, 18),
            "Lower gradient step",
        )
    )
    custom = workbook.create_sheet("Custom rescues")
    custom.append(("Keep",))
    custom.append(("unchanged",))
    workbook.save(path)


def test_legacy_proposals_migrate_to_rescue_request_rows(tmp_path: Path) -> None:
    path = tmp_path / "registration_evaluation.xlsx"
    _write_legacy_workbook(path)

    data = read_evaluation_workbook(path)

    assert len(data.run_evaluations) == 1
    assert data.run_evaluations[0].run_path == "registration_runs/baseline"
    assert data.run_evaluations[0].comments == "Keep this note"
    assert data.rescue_requests[0].request.subject_id == "A"
    assert data.rescue_requests[0].request.strategy == "Lower gradient step"
    assert data.rescue_requests[0].request.source_run == "registration_runs/baseline"


def test_plan_apply_preserves_human_fields_and_auxiliary_sheets(
    tmp_path: Path,
) -> None:
    path = tmp_path / "registration_evaluation.xlsx"
    _write_legacy_workbook(path)
    strategies = (RescueStrategy("lower_gradient", "Lower gradient step"),)
    plan = build_evaluation_workbook_plan(
        path=path,
        discovered_runs=(
            RunEvaluation(
                "A",
                "registration_runs/baseline",
                preset_name="updated_preset",
                manifest_status="success",
            ),
            RunEvaluation(
                "B",
                "registration_runs/baseline",
                preset_name="baseline",
                manifest_status="planned",
            ),
        ),
        strategies=strategies,
        evaluator="Evaluator",
        rescue_statuses={
            RescueRequestKey("A", "lower_gradient"): RescueRequestStatus.GENERATED
        },
    )

    assert plan.added_run_count == 1
    assert plan.migrated_rescue_count == 1
    apply_evaluation_workbook_plan(plan)

    assert plan.backup_path.is_file()
    workbook = load_workbook(path)
    assert "Custom rescues" not in workbook.sheetnames
    assert RESCUE_REQUESTS_SHEET in workbook.sheetnames
    assert WORKFLOW_LISTS_SHEET in workbook.sheetnames
    assert workbook[WORKFLOW_LISTS_SHEET].sheet_state == "hidden"
    run_headers = [cell.value for cell in workbook[RUN_EVALUATIONS_SHEET][1]]
    assert "Proposed rescue" not in run_headers
    assert len(workbook[RESCUE_REQUESTS_SHEET].data_validations.dataValidation) == 3
    workbook.close()

    synchronized = read_evaluation_workbook(path)
    first = synchronized.run_evaluations[0]
    assert first.preset_name == "updated_preset"
    assert first.manifest_status == "success"
    assert first.result == "minor_localized"
    assert first.comments == "Keep this note"
    assert first.selected
    assert first.evaluator == "Evaluator"
    assert first.evaluation_date.date() == date(2026, 9, 18)
    assert len(synchronized.rescue_requests) == 1
    assert synchronized.rescue_requests[0].status == "generated"

    second_plan = build_evaluation_workbook_plan(
        path=path,
        discovered_runs=synchronized.run_evaluations,
        strategies=strategies,
        evaluator="Evaluator",
    )
    assert second_plan.migrated_rescue_count == 0
    assert len(second_plan.data.rescue_requests) == 1


@pytest.mark.parametrize(
    ("statuses", "results", "expected"),
    (
        (("success",) * 3, ("overwarping",) * 3, "completed"),
        (("success",) * 3, ("overwarping", "", "overwarping"), "generated"),
        (("failed",) * 3, ("",) * 3, "failed"),
        (("success", "failed", "planned"), ("overwarping", "", ""), "generated"),
    ),
)
def test_refresh_closes_only_finished_rescue_sweeps(
    tmp_path: Path, statuses, results, expected: str
) -> None:
    path = tmp_path / "registration_evaluation.xlsx"
    workbook = Workbook()
    run_sheet = workbook.active
    run_sheet.title = RUN_EVALUATIONS_SHEET
    run_sheet.append(RUN_EVALUATION_COLUMNS)
    variants = ("gs0p01", "gs0p02", "gs0p025")
    runs = tuple(
        RunEvaluation(
            "A", f"registration_runs/lower_gradient_step_rescue/{variant}",
            manifest_status=status, result=result,
        )
        for variant, status, result in zip(variants, statuses, results)
    )
    for run in runs:
        run_sheet.append((run.subject_id, run.run_path, "", run.manifest_status,
                          run.result, "", "", "", None))
    rescue_sheet = workbook.create_sheet(RESCUE_REQUESTS_SHEET)
    rescue_sheet.append(RESCUE_REQUEST_COLUMNS)
    rescue_sheet.append(("A", "lower gradient step", "", "", None, None, None,
                         "", "generated"))
    workbook.save(path)

    plan = build_evaluation_workbook_plan(
        path=path,
        discovered_runs=runs,
        strategies=(RescueStrategy("lower_gradient_step", "lower gradient step"),),
    )

    assert plan.data.rescue_requests[0].status == expected


@pytest.mark.parametrize(
    ("manifest_status", "result", "run_path", "expected_status"),
    (
        ("success", "excellent", "registration_runs/masking_rescue/from_padding_rescue_wr20um_pad500um", "completed"),
        ("success", "", "registration_runs/masking_rescue/from_padding_rescue_wr20um_pad500um", "generated"),
        ("failed", "", "registration_runs/masking_rescue/from_padding_rescue_wr20um_pad500um", "failed"),
        ("success", "excellent", "registration_runs/masking_rescue/from_other_run", "generated"),
    ),
)
def test_refresh_closes_evaluated_masking_rescue(
    tmp_path: Path, manifest_status: str, result: str, run_path: str,
    expected_status: str,
) -> None:
    source_run = "registration_runs/padding_rescue/wr20um_pad500um"
    path = tmp_path / "registration_evaluation.xlsx"
    workbook = Workbook()
    run_sheet = workbook.active
    run_sheet.title = RUN_EVALUATIONS_SHEET
    run_sheet.append(RUN_EVALUATION_COLUMNS)
    run_sheet.append(("IEB0167", run_path, "", manifest_status, result,
                      "", "yes", "", None))
    rescue_sheet = workbook.create_sheet(RESCUE_REQUESTS_SHEET)
    rescue_sheet.append(RESCUE_REQUEST_COLUMNS)
    rescue_sheet.append(("IEB0167", "masking", source_run, "", None, None,
                         None, "", "generated"))
    workbook.save(path)

    masking = next(
        strategy for strategy in canonical_rescue_strategies()
        if strategy.strategy_id == "masking"
    )
    assert expected_rescue_run_paths(
        RescueRequest("IEB0167", "masking", source_run), masking
    ) == ("registration_runs/masking_rescue/from_padding_rescue_wr20um_pad500um",)
    plan = build_evaluation_workbook_plan(
        path=path,
        discovered_runs=(RunEvaluation("IEB0167", run_path,
                                       manifest_status=manifest_status),),
        strategies=canonical_rescue_strategies(),
    )

    assert plan.data.rescue_requests[0].status == expected_status
    apply_evaluation_workbook_plan(plan)
    assert read_evaluation_workbook(path).rescue_requests[0].status == expected_status


def test_apply_rejects_workbook_changed_after_planning(tmp_path: Path) -> None:
    path = tmp_path / "registration_evaluation.xlsx"
    _write_legacy_workbook(path)
    plan = build_evaluation_workbook_plan(
        path=path,
        discovered_runs=read_evaluation_workbook(path).run_evaluations,
        strategies=(RescueStrategy("lower_gradient", "Lower gradient step"),),
    )
    with path.open("ab") as handle:
        handle.write(b"changed")

    with pytest.raises(RuntimeError, match="changed after planning"):
        apply_evaluation_workbook_plan(plan)


def test_rescue_request_status_must_be_derived_vocabulary() -> None:
    with pytest.raises(ValueError, match="Unsupported derived"):
        RescueRequestRecord(
            request=RescueRequest("A", "Lower gradient step"),
            status="manually done",
        )


def test_auxiliary_sheet_is_created_once_and_then_preserved(tmp_path: Path) -> None:
    path = tmp_path / "registration_evaluation.xlsx"
    plan = build_evaluation_workbook_plan(
        path=path,
        discovered_runs=(RunEvaluation("A", "registration_runs/baseline"),),
        strategies=(),
        auxiliary_sheet_headers={"Review notes": ("Request ID", "Ready")},
    )
    apply_evaluation_workbook_plan(plan)
    workbook = load_workbook(path)
    workbook["Review notes"].append(("request_1", "yes"))
    workbook.save(path)
    workbook.close()

    rerun = build_evaluation_workbook_plan(
        path=path,
        discovered_runs=read_evaluation_workbook(path).run_evaluations,
        strategies=(),
        auxiliary_sheet_headers={"Review notes": ("Request ID", "Ready")},
    )
    apply_evaluation_workbook_plan(rerun)
    workbook = load_workbook(path)
    assert workbook["Review notes"]["A2"].value == "request_1"
    workbook.close()


def test_legacy_custom_rows_migrate_into_unified_rescue_requests(tmp_path: Path) -> None:
    path = tmp_path / "registration_evaluation.xlsx"
    workbook = Workbook()
    runs = workbook.active
    runs.title = RUN_EVALUATIONS_SHEET
    runs.append((*RUN_EVALUATION_COLUMNS, "Proposed rescue"))
    runs.append(
        (
            "A",
            "registration_runs/baseline",
            "baseline",
            "success",
            "underwarping",
            "reviewed",
            "",
            "Evaluator",
            None,
            "custom",
        )
    )
    custom = workbook.create_sheet("Custom rescues")
    custom.append(
        (
            "Request ID",
            "Subject ID",
            "Variant name",
            "Gradient step",
            "Working resolution (um)",
            "Padding (um)",
            "Rationale",
            "Ready",
        )
    )
    custom.append(("test", "A", "gs0p08", 0.08, None, None, "comparison", "yes"))
    workbook.save(path)

    data = read_evaluation_workbook(path)

    assert len(data.rescue_requests) == 1
    request = data.rescue_requests[0].request
    assert request.strategy == "custom"
    assert request.source_run == "registration_runs/baseline"
    assert request.variant_name == "gs0p08"
    assert request.gradient_step == 0.08
    assert request.notes == "comparison"
