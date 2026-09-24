"""Read and plan synchronization of registration evaluation workbooks."""

from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path
from typing import Iterable, Mapping

from openpyxl import load_workbook

from .decisions import WorkflowContractError, make_rescue_request_key, rescue_strategy_lookup
from .rescue_catalog import expected_rescue_run_paths
from .models import (
    EvaluationWorkbookData,
    RescueRequest,
    RescueRequestKey,
    RescueRequestRecord,
    RescueRequestStatus,
    RescueStrategy,
    RunEvaluation,
    normalize_run_path,
    normalize_subject_id,
    normalize_text,
)
from .schema import (
    LEGACY_RESCUE_COLUMN,
    RESCUE_REQUESTS_SHEET,
    RUN_EVALUATIONS_SHEET,
)


LEGACY_CUSTOM_RESCUES_SHEET = "Custom rescues"


@dataclass(frozen=True)
class EvaluationWorkbookPlan:
    path: Path
    data: EvaluationWorkbookData
    strategies: tuple[RescueStrategy, ...]
    source_sha256: str | None
    existing_run_count: int
    added_run_count: int
    migrated_rescue_count: int
    auxiliary_sheet_headers: tuple[tuple[str, tuple[str, ...]], ...] = ()

    @property
    def backup_path(self) -> Path:
        return self.path.with_name(f"{self.path.stem}_previous{self.path.suffix}")


def file_sha256(path: Path) -> str | None:
    return sha256(path.read_bytes()).hexdigest() if path.is_file() else None


def read_evaluation_workbook(path: Path | str) -> EvaluationWorkbookData:
    """Read canonical sheets and migrate legacy proposals in memory."""

    path = Path(path)
    if not path.is_file():
        return EvaluationWorkbookData((), ())
    workbook = load_workbook(path, data_only=False)
    try:
        return _read_loaded_workbook(workbook, path)
    finally:
        workbook.close()


def _read_loaded_workbook(workbook, path: Path) -> EvaluationWorkbookData:
    if RUN_EVALUATIONS_SHEET not in workbook.sheetnames:
        raise WorkflowContractError(
            f"Evaluation workbook has no {RUN_EVALUATIONS_SHEET!r} sheet: {path}"
        )
    run_sheet = workbook[RUN_EVALUATIONS_SHEET]
    headers = _header_map(run_sheet)
    missing = sorted({"Subject ID", "Run path"}.difference(headers))
    if missing:
        raise WorkflowContractError(
            f"Run evaluations sheet is missing columns: {missing}"
        )

    run_rows: list[RunEvaluation] = []
    migrated_requests: list[RescueRequestRecord] = []
    for row_number in range(2, run_sheet.max_row + 1):
        subject_value = _cell_value(run_sheet, headers, row_number, "Subject ID")
        run_value = _cell_value(run_sheet, headers, row_number, "Run path")
        if not normalize_text(subject_value) and not normalize_text(run_value):
            continue
        run = RunEvaluation(
            subject_id=subject_value,
            run_path=run_value,
            preset_name=_cell_value(run_sheet, headers, row_number, "Preset name"),
            manifest_status=_cell_value(run_sheet, headers, row_number, "Manifest status"),
            result=_cell_value(run_sheet, headers, row_number, "Result"),
            comments=_cell_value(run_sheet, headers, row_number, "Comments"),
            selected=_cell_value(run_sheet, headers, row_number, "Selected"),
            evaluator=_cell_value(run_sheet, headers, row_number, "Evaluator"),
            evaluation_date=_cell_value(run_sheet, headers, row_number, "Evaluation date"),
        )
        run_rows.append(run)
        legacy_rescue = normalize_text(
            _cell_value(run_sheet, headers, row_number, LEGACY_RESCUE_COLUMN)
        )
        if legacy_rescue:
            migrated_requests.append(
                RescueRequestRecord(
                    RescueRequest(
                        subject_id=run.subject_id,
                        strategy=legacy_rescue,
                        source_run=run.run_path,
                    )
                )
            )

    _raise_for_duplicate_runs(run_rows)
    rescue_rows = _read_rescue_requests(workbook)
    rescue_rows.extend(migrated_requests)
    legacy_custom_rows = _read_legacy_custom_requests(workbook, rescue_rows)
    if legacy_custom_rows:
        migrated_subjects = {
            record.request.subject_id for record in legacy_custom_rows
        }
        rescue_rows = [
            record
            for record in rescue_rows
            if not (
                record.request.subject_id in migrated_subjects
                and record.request.strategy.casefold() == "custom"
            )
        ]
        rescue_rows.extend(legacy_custom_rows)
    return EvaluationWorkbookData(tuple(run_rows), tuple(rescue_rows))


def build_evaluation_workbook_plan(
    *,
    path: Path | str,
    discovered_runs: Iterable[RunEvaluation],
    strategies: Iterable[RescueStrategy],
    evaluator: str = "",
    rescue_statuses: Mapping[RescueRequestKey, RescueRequestStatus | str] | None = None,
    auxiliary_sheet_headers: Mapping[str, Iterable[str]] | None = None,
) -> EvaluationWorkbookPlan:
    """Plan a non-destructive workbook synchronization."""

    path = Path(path)
    existing = read_evaluation_workbook(path)
    discovered = tuple(discovered_runs)
    _raise_for_duplicate_runs(discovered)
    strategy_tuple = tuple(strategies)
    rescue_strategy_lookup(strategy_tuple)

    existing_by_key = {
        (row.subject_id, row.run_path): row for row in existing.run_evaluations
    }
    discovered_by_key = {(row.subject_id, row.run_path): row for row in discovered}
    retained_existing = {
        key
        for key, row in existing_by_key.items()
        if key in discovered_by_key or _has_human_evaluation(row)
    }
    synchronized_runs: list[RunEvaluation] = []
    for key in sorted(retained_existing | set(discovered_by_key)):
        old = existing_by_key.get(key)
        current = discovered_by_key.get(key)
        if current is None:
            assert old is not None
            synchronized_runs.append(old)
            continue
        synchronized_runs.append(
            RunEvaluation(
                subject_id=current.subject_id,
                run_path=current.run_path,
                preset_name=current.preset_name,
                manifest_status=current.manifest_status,
                result=old.result if old is not None else "",
                comments=old.comments if old is not None else "",
                selected=old.selected if old is not None else False,
                evaluator=(
                    old.evaluator
                    if old is not None and old.evaluator
                    else normalize_text(evaluator)
                ),
                evaluation_date=old.evaluation_date if old is not None else None,
            )
        )

    synchronized_requests = _synchronize_rescue_statuses(
        existing.rescue_requests,
        strategies=strategy_tuple,
        rescue_statuses=rescue_statuses or {},
        run_evaluations=synchronized_runs,
    )
    existing_keys = set(existing_by_key)
    return EvaluationWorkbookPlan(
        path=path,
        data=EvaluationWorkbookData(
            tuple(synchronized_runs), tuple(synchronized_requests)
        ),
        strategies=strategy_tuple,
        source_sha256=file_sha256(path),
        existing_run_count=len(existing.run_evaluations),
        added_run_count=len(set(discovered_by_key) - existing_keys),
        migrated_rescue_count=(
            len(existing.rescue_requests)
            if path.is_file() and not _workbook_has_sheet(path, RESCUE_REQUESTS_SHEET)
            else 0
        ),
        auxiliary_sheet_headers=tuple(
            (str(title), tuple(str(column) for column in columns))
            for title, columns in (auxiliary_sheet_headers or {}).items()
        ),
    )


def summarize_evaluation_workbook_plan(plan: EvaluationWorkbookPlan) -> str:
    return "\n".join(
        [
            f"Registration evaluation workbook plan: {plan.path}",
            f"Existing run rows: {plan.existing_run_count}",
            f"Synchronized run rows: {len(plan.data.run_evaluations)}",
            f"New run rows: {plan.added_run_count}",
            f"Rescue request rows: {len(plan.data.rescue_requests)}",
            f"Legacy rescue rows migrated: {plan.migrated_rescue_count}",
        ]
    )


def _read_rescue_requests(workbook) -> list[RescueRequestRecord]:
    if RESCUE_REQUESTS_SHEET not in workbook.sheetnames:
        return []
    sheet = workbook[RESCUE_REQUESTS_SHEET]
    headers = _header_map(sheet)
    required = {"Subject ID", "Rescue strategy"}
    missing = sorted(required.difference(headers))
    if missing:
        raise WorkflowContractError(
            f"Rescue requests sheet is missing columns: {missing}"
        )
    records: list[RescueRequestRecord] = []
    for row_number in range(2, sheet.max_row + 1):
        subject = _cell_value(sheet, headers, row_number, "Subject ID")
        strategy = _cell_value(sheet, headers, row_number, "Rescue strategy")
        if not normalize_text(subject) and not normalize_text(strategy):
            continue
        records.append(
            RescueRequestRecord(
                RescueRequest(
                    subject_id=subject,
                    strategy=strategy,
                    source_run=_cell_value(sheet, headers, row_number, "Source run"),
                    variant_name=_cell_value(sheet, headers, row_number, "Variant name"),
                    gradient_step=_cell_value(sheet, headers, row_number, "Gradient step"),
                    working_resolution_um=_cell_value(
                        sheet, headers, row_number, "Working resolution (um)"
                    ),
                    padding_um=_cell_value(sheet, headers, row_number, "Padding (um)"),
                    notes=_cell_value(sheet, headers, row_number, "Notes"),
                ),
                status=_cell_value(sheet, headers, row_number, "Status"),
            )
        )
    return records


def _read_legacy_custom_requests(
    workbook, existing: Iterable[RescueRequestRecord]
) -> list[RescueRequestRecord]:
    if LEGACY_CUSTOM_RESCUES_SHEET not in workbook.sheetnames:
        return []
    sheet = workbook[LEGACY_CUSTOM_RESCUES_SHEET]
    headers = _header_map(sheet)
    required = {
        "Subject ID", "Variant name", "Gradient step", "Working resolution (um)",
        "Padding (um)", "Rationale", "Ready",
    }
    if not required.issubset(headers):
        return []
    source_runs = {
        record.request.subject_id: record.request.source_run
        for record in existing
        if record.request.strategy.casefold() == "custom"
    }
    records: list[RescueRequestRecord] = []
    for row_number in range(2, sheet.max_row + 1):
        if normalize_text(_cell_value(sheet, headers, row_number, "Ready")).casefold() not in {
            "1", "true", "yes"
        }:
            continue
        subject_id = normalize_subject_id(
            _cell_value(sheet, headers, row_number, "Subject ID")
        )
        records.append(
            RescueRequestRecord(
                RescueRequest(
                    subject_id=subject_id,
                    strategy="custom",
                    source_run=source_runs.get(subject_id, ""),
                    variant_name=_cell_value(sheet, headers, row_number, "Variant name"),
                    gradient_step=_cell_value(sheet, headers, row_number, "Gradient step"),
                    working_resolution_um=_cell_value(
                        sheet, headers, row_number, "Working resolution (um)"
                    ),
                    padding_um=_cell_value(sheet, headers, row_number, "Padding (um)"),
                    notes=_cell_value(sheet, headers, row_number, "Rationale"),
                )
            )
        )
    return records


def _synchronize_rescue_statuses(
    records: Iterable[RescueRequestRecord],
    *,
    strategies: tuple[RescueStrategy, ...],
    rescue_statuses: Mapping[RescueRequestKey, RescueRequestStatus | str],
    run_evaluations: Iterable[RunEvaluation],
) -> list[RescueRequestRecord]:
    lookup = rescue_strategy_lookup(strategies)
    runs_by_key = {
        (row.subject_id, row.run_path): row for row in run_evaluations
    }
    synchronized: list[RescueRequestRecord] = []
    for record in records:
        strategy = lookup.get(record.request.strategy.casefold())
        status = record.status
        if strategy is not None:
            key = make_rescue_request_key(record.request, strategy)
            if key in rescue_statuses:
                status = RescueRequestStatus(rescue_statuses[key]).value
            elif status == RescueRequestStatus.GENERATED.value:
                expected = expected_rescue_run_paths(record.request, strategy)
                runs = [
                    runs_by_key.get((record.request.subject_id, path))
                    for path in expected
                ]
                if expected and all(
                    run is not None
                    and (
                        run.manifest_status == "failed"
                        or (run.manifest_status == "success" and run.result)
                    )
                    for run in runs
                ):
                    status = (
                        RescueRequestStatus.COMPLETED.value
                        if any(run.manifest_status == "success" for run in runs)
                        else RescueRequestStatus.FAILED.value
                    )
        synchronized.append(RescueRequestRecord(record.request, status))
    return synchronized


def _header_map(sheet) -> dict[str, int]:
    headers: dict[str, int] = {}
    duplicates: list[str] = []
    for column in range(1, sheet.max_column + 1):
        name = normalize_text(sheet.cell(1, column).value)
        if not name:
            continue
        if name in headers:
            duplicates.append(name)
        headers[name] = column
    if duplicates:
        raise WorkflowContractError(
            f"Sheet {sheet.title!r} has duplicate columns: {sorted(set(duplicates))}"
        )
    return headers


def _cell_value(sheet, headers: Mapping[str, int], row: int, column: str):
    index = headers.get(column)
    return sheet.cell(row, index).value if index is not None else None


def _raise_for_duplicate_runs(rows: Iterable[RunEvaluation]) -> None:
    seen: set[tuple[str, str]] = set()
    duplicates: set[tuple[str, str]] = set()
    for row in rows:
        key = (row.subject_id, row.run_path)
        if key in seen:
            duplicates.add(key)
        seen.add(key)
    if duplicates:
        raise WorkflowContractError(
            f"Run evaluations contain duplicate subject/run rows: {sorted(duplicates)}"
        )


def _has_human_evaluation(row: RunEvaluation) -> bool:
    return any((row.result, row.comments, row.selected, row.evaluation_date is not None))


def _workbook_has_sheet(path: Path, sheet_name: str) -> bool:
    workbook = load_workbook(path, read_only=True)
    try:
        return sheet_name in workbook.sheetnames
    finally:
        workbook.close()


__all__ = [
    "EvaluationWorkbookPlan",
    "build_evaluation_workbook_plan",
    "file_sha256",
    "read_evaluation_workbook",
    "summarize_evaluation_workbook_plan",
]
