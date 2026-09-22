"""Render and atomically write registration evaluation workbooks."""

from __future__ import annotations

from copy import copy
from pathlib import Path
import shutil
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from .decisions import make_rescue_request_key
from .models import EvaluationWorkbookData, RescueRequestStatus, RescueStrategy, normalize_text
from .schema import (
    CATEGORIES_SHEET,
    RESCUE_REQUEST_COLUMNS,
    RESCUE_REQUESTS_SHEET,
    RESULT_CATEGORIES,
    RUN_EVALUATION_COLUMNS,
    RUN_EVALUATIONS_SHEET,
    WORKFLOW_LISTS_SHEET,
)
from .workbook_reader import EvaluationWorkbookPlan, file_sha256


LEGACY_CUSTOM_RESCUES_SHEET = "Custom rescues"


def apply_evaluation_workbook_plan(plan: EvaluationWorkbookPlan) -> Path:
    """Atomically apply a plan after verifying its source workbook is unchanged."""

    if file_sha256(plan.path) != plan.source_sha256:
        raise RuntimeError(
            f"Evaluation workbook changed after planning and was not rewritten: {plan.path}"
        )
    plan.path.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(plan.path) if plan.path.is_file() else Workbook()
    if not plan.path.is_file() and workbook.active is not None:
        workbook.remove(workbook.active)

    _replace_run_sheet(workbook, plan.data.run_evaluations)
    _replace_rescue_sheet(workbook, plan.data.rescue_requests)
    _replace_categories_sheet(workbook, plan.strategies)
    _replace_lists_sheet(workbook, plan.data, plan.strategies)
    _ensure_auxiliary_sheets(workbook, plan.auxiliary_sheet_headers)
    if LEGACY_CUSTOM_RESCUES_SHEET in workbook.sheetnames:
        workbook.remove(workbook[LEGACY_CUSTOM_RESCUES_SHEET])
    _add_validations(workbook, plan.data, plan.strategies)

    handle = tempfile.NamedTemporaryFile(
        prefix=f".{plan.path.stem}_",
        suffix=plan.path.suffix,
        dir=plan.path.parent,
        delete=False,
    )
    temporary = Path(handle.name)
    handle.close()
    try:
        try:
            workbook.save(temporary)
        finally:
            workbook.close()
        if plan.path.is_file():
            _replace_backup(plan.path, plan.backup_path)
        temporary.replace(plan.path)
    finally:
        temporary.unlink(missing_ok=True)
    return plan.path


def _replace_run_sheet(workbook, rows) -> None:
    sheet = _replace_sheet(workbook, RUN_EVALUATIONS_SHEET)
    sheet.append(RUN_EVALUATION_COLUMNS)
    for row in rows:
        sheet.append(
            (
                row.subject_id, row.run_path, row.preset_name, row.manifest_status,
                row.result, row.comments, "yes" if row.selected else "",
                row.evaluator, row.evaluation_date,
            )
        )
    _format_table_sheet(sheet)


def _replace_rescue_sheet(workbook, rows) -> None:
    sheet = _replace_sheet(workbook, RESCUE_REQUESTS_SHEET)
    sheet.append(RESCUE_REQUEST_COLUMNS)
    for record in rows:
        request = record.request
        sheet.append(
            (
                request.subject_id, request.strategy, request.source_run,
                request.variant_name, request.gradient_step,
                request.working_resolution_um, request.padding_um, request.notes,
                record.status,
            )
        )
    _format_table_sheet(sheet)


def _replace_categories_sheet(workbook, strategies: tuple[RescueStrategy, ...]) -> None:
    sheet = _replace_sheet(workbook, CATEGORIES_SHEET)
    sheet.append(("Entry type", "Value", "Definition"))
    for value, definition in RESULT_CATEGORIES.items():
        sheet.append(("Result category", value, definition))
    sheet.append(("Selection marker", "yes", "The final selected run or exclusion."))
    sheet.append(("Selection marker", "no", "An evaluated, unselected run."))
    for strategy in strategies:
        sheet.append(
            (
                "Rescue strategy",
                strategy.workbook_label,
                f"{strategy.kind.value}; source run is {strategy.source_run_requirement.value}",
            )
        )
    for status in RescueRequestStatus:
        sheet.append(("Derived request status", status.value, "Set by synchronization."))
    _format_table_sheet(sheet)


def _replace_lists_sheet(
    workbook, data: EvaluationWorkbookData, strategies: tuple[RescueStrategy, ...]
) -> None:
    sheet = _replace_sheet(workbook, WORKFLOW_LISTS_SHEET)
    sheet.append(
        ("Results", "Selected markers", "Rescue strategies", "Subjects", "", "Run subject", "Run path")
    )
    subjects = sorted({row.subject_id for row in data.run_evaluations})
    runs = sorted((row.subject_id, row.run_path) for row in data.run_evaluations)
    maximum = max(len(RESULT_CATEGORIES), 2, len(strategies), len(subjects), len(runs), 1)
    results = list(RESULT_CATEGORIES)
    selections = ["yes", "no"]
    strategy_labels = [strategy.workbook_label for strategy in strategies]
    for index in range(maximum):
        run_subject, run_path = runs[index] if index < len(runs) else ("", "")
        sheet.append(
            (
                results[index] if index < len(results) else "",
                selections[index] if index < len(selections) else "",
                strategy_labels[index] if index < len(strategy_labels) else "",
                subjects[index] if index < len(subjects) else "",
                "",
                run_subject,
                run_path,
            )
        )
    sheet.sheet_state = "hidden"


def _ensure_auxiliary_sheets(workbook, sheets) -> None:
    protected = {
        RUN_EVALUATIONS_SHEET,
        RESCUE_REQUESTS_SHEET,
        CATEGORIES_SHEET,
        WORKFLOW_LISTS_SHEET,
    }
    for title, headers in sheets:
        if title in protected:
            raise ValueError(f"Auxiliary sheet conflicts with a canonical sheet: {title}")
        if title in workbook.sheetnames:
            continue
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        _format_table_sheet(sheet)


def _add_validations(workbook, data: EvaluationWorkbookData, strategies) -> None:
    run_sheet = workbook[RUN_EVALUATIONS_SHEET]
    rescue_sheet = workbook[RESCUE_REQUESTS_SHEET]
    subject_count = len({row.subject_id for row in data.run_evaluations})
    _add_list_validation(
        run_sheet,
        RUN_EVALUATION_COLUMNS.index("Result") + 1,
        f"'{WORKFLOW_LISTS_SHEET}'!$A$2:$A${len(RESULT_CATEGORIES) + 1}",
    )
    _add_list_validation(
        run_sheet,
        RUN_EVALUATION_COLUMNS.index("Selected") + 1,
        f"'{WORKFLOW_LISTS_SHEET}'!$B$2:$B$3",
    )
    if subject_count:
        _add_list_validation(
            rescue_sheet,
            RESCUE_REQUEST_COLUMNS.index("Subject ID") + 1,
            f"'{WORKFLOW_LISTS_SHEET}'!$D$2:$D${subject_count + 1}",
        )
    if strategies:
        _add_list_validation(
            rescue_sheet,
            RESCUE_REQUEST_COLUMNS.index("Rescue strategy") + 1,
            f"'{WORKFLOW_LISTS_SHEET}'!$C$2:$C${len(strategies) + 1}",
        )
    source_column = RESCUE_REQUEST_COLUMNS.index("Source run") + 1
    source_letter = rescue_sheet.cell(1, source_column).column_letter
    validation = DataValidation(
        type="list",
        formula1=(
            f"=OFFSET('{WORKFLOW_LISTS_SHEET}'!$G$1,"
            f"MATCH($A2,'{WORKFLOW_LISTS_SHEET}'!$F:$F,0)-1,0,"
            f"COUNTIF('{WORKFLOW_LISTS_SHEET}'!$F:$F,$A2),1)"
        ),
        allow_blank=True,
        errorStyle="stop",
    )
    rescue_sheet.add_data_validation(validation)
    validation.add(f"{source_letter}2:{source_letter}5000")


def _add_list_validation(sheet, column: int, formula: str) -> None:
    validation = DataValidation(type="list", formula1=formula, allow_blank=True, errorStyle="stop")
    sheet.add_data_validation(validation)
    letter = sheet.cell(1, column).column_letter
    validation.add(f"{letter}2:{letter}5000")


def _replace_sheet(workbook, title: str):
    index = len(workbook.worksheets)
    if title in workbook.sheetnames:
        old = workbook[title]
        index = workbook.worksheets.index(old)
        workbook.remove(old)
    return workbook.create_sheet(title, index)


def _format_table_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        font = copy(cell.font)
        font.bold = True
        cell.font = font
    for column_cells in sheet.columns:
        values = [normalize_text(cell.value) for cell in column_cells]
        width = min(max((len(value) for value in values), default=0) + 2, 60)
        sheet.column_dimensions[column_cells[0].column_letter].width = max(width, 12)


def _replace_backup(source: Path, destination: Path) -> None:
    handle = tempfile.NamedTemporaryFile(
        prefix=f".{destination.stem}_",
        suffix=destination.suffix,
        dir=destination.parent,
        delete=False,
    )
    temporary = Path(handle.name)
    handle.close()
    try:
        shutil.copy2(source, temporary)
        temporary.replace(destination)
    finally:
        temporary.unlink(missing_ok=True)


__all__ = ["apply_evaluation_workbook_plan"]
